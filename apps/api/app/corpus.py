"""
In-memory institutional corpus index.

Sources:
1. data/processed/*.json (RunPulse structured extracts)
2. Work Orders/**/*.docx (python-docx)
3. Work Orders/**/*.pdf (pypdf text layer when present)
"""
from __future__ import annotations

import json
import math
import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[3]
PROCESSED = ROOT / "data" / "processed"
WORK_ORDERS = ROOT / "Work Orders"

_TOKEN = re.compile(r"[a-z0-9][a-z0-9\-/&]{1,}", re.I)
_STOP = {
    "the", "and", "for", "with", "from", "that", "this", "are", "was", "were",
    "of", "to", "in", "on", "at", "by", "an", "or", "as", "be", "is", "it",
    "a", "not", "will", "shall", "have", "has", "had", "been", "their",
}


def _tokens(text: str) -> list[str]:
    return [t.lower() for t in _TOKEN.findall(text or "") if t.lower() not in _STOP and len(t) > 2]


@dataclass
class Chunk:
    chunk_id: str
    doc_id: str
    title: str
    ministry: str
    date: str | None
    domains: list[str]
    text: str
    source: str
    value_inr: float = 0.0
    kind: str = "processed_json"
    token_set: set[str] = field(default_factory=set)


@dataclass
class CorpusDoc:
    doc_id: str
    title: str
    ministry: str
    date: str | None
    domains: list[str]
    deliverables: str
    full_text: str
    value_inr: float
    path: str
    kind: str


def _split_passages(text: str, size: int = 700, overlap: int = 120) -> list[str]:
    text = re.sub(r"\s+", " ", (text or "").strip())
    if not text:
        return []
    if len(text) <= size:
        return [text]
    out: list[str] = []
    i = 0
    while i < len(text):
        out.append(text[i : i + size])
        i += max(size - overlap, 1)
    return out


def _rel(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT)).replace("\\", "/")
    except ValueError:
        return str(path).replace("\\", "/")


def _title_from_name(name: str) -> str:
    return re.sub(r"[_\-]+", " ", Path(name).stem).strip()


def _guess_domains(name: str, text: str) -> list[str]:
    blob = f"{name} {text}".lower()
    tags: list[str] = []
    mapping = [
        ("cpgrams", "CPGRAMS"),
        ("nesda", "NeSDA"),
        ("scdpm", "SCDPM"),
        ("sanitation", "Sanitation"),
        ("impact", "Impact Assessment"),
        ("dpiit", "DPIIT"),
        ("doppw", "DoPPW"),
        ("media", "Media"),
        ("iedm", "IEDM"),
        ("darpg", "DARPG"),
        ("pmu", "PMU"),
    ]
    for key, label in mapping:
        if key in blob:
            tags.append(label)
    return tags or ["General"]


def _extract_docx(path: Path) -> str:
    try:
        import docx  # type: ignore
    except ImportError:
        return ""
    try:
        d = docx.Document(str(path))
        parts = [p.text.strip() for p in d.paragraphs if p.text and p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception:
        return ""


def _extract_pdf(path: Path, max_pages: int = 40) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        return ""
    try:
        reader = PdfReader(str(path))
        parts: list[str] = []
        for page in reader.pages[:max_pages]:
            try:
                t = (page.extract_text() or "").strip()
            except Exception:
                t = ""
            if t:
                parts.append(t)
        return "\n".join(parts)
    except Exception:
        return ""


def _add_doc(
    docs: list[CorpusDoc],
    chunks: list[Chunk],
    *,
    doc_id: str,
    title: str,
    ministry: str,
    date: str | None,
    domains: list[str],
    deliverables: str,
    full_text: str,
    value_inr: float,
    path: str,
    kind: str,
) -> None:
    existing = next((d for d in docs if d.doc_id == doc_id), None)
    if existing and existing.kind == "processed_json" and kind != "processed_json":
        return
    if existing and existing.kind != "processed_json" and kind == "processed_json":
        docs[:] = [d for d in docs if d.doc_id != doc_id]
        chunks[:] = [c for c in chunks if c.doc_id != doc_id]
    elif existing:
        return

    docs.append(
        CorpusDoc(
            doc_id=doc_id,
            title=title,
            ministry=ministry,
            date=date,
            domains=domains,
            deliverables=deliverables,
            full_text=full_text,
            value_inr=value_inr,
            path=path,
            kind=kind,
        )
    )
    base = f"{title}. {deliverables}. {full_text}".strip()
    passages = _split_passages(base) or [title]
    for i, passage in enumerate(passages):
        chunks.append(
            Chunk(
                chunk_id=f"{doc_id}::c{i}",
                doc_id=doc_id,
                title=title,
                ministry=ministry,
                date=date,
                domains=domains,
                text=passage,
                source=path,
                value_inr=value_inr,
                kind=kind,
                token_set=set(_tokens(passage)),
            )
        )


def _load_processed(docs: list[CorpusDoc], chunks: list[Chunk]) -> None:
    if not PROCESSED.exists():
        return
    for path in sorted(PROCESSED.glob("*.json")):
        try:
            raw = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        meta = raw.get("meta") or {}
        content = raw.get("content") or {}
        doc_id = str(raw.get("doc_id") or meta.get("doc_id") or path.stem)
        try:
            value_inr = float(meta.get("value_inr") or 0)
        except (TypeError, ValueError):
            value_inr = 0.0
        _add_doc(
            docs,
            chunks,
            doc_id=doc_id,
            title=str(meta.get("project_subject") or doc_id),
            ministry=str(meta.get("ministry") or "Unknown"),
            date=meta.get("date"),
            domains=list(meta.get("domains") or []),
            deliverables=str(meta.get("deliverables") or ""),
            full_text=str(content.get("full_text") or ""),
            value_inr=value_inr,
            path=_rel(path),
            kind="processed_json",
        )


def _load_work_orders(docs: list[CorpusDoc], chunks: list[Chunk]) -> dict[str, int]:
    stats = {"docx": 0, "pdf": 0, "pdf_empty_text": 0}
    if not WORK_ORDERS.exists():
        return stats
    for path in sorted(WORK_ORDERS.rglob("*")):
        if not path.is_file() or path.name.startswith("~$"):
            continue
        ext = path.suffix.lower()
        if ext not in {".docx", ".pdf"}:
            continue
        title = _title_from_name(path.name)
        doc_id = path.stem
        if ext == ".docx":
            text = _extract_docx(path)
            kind = "work_order_docx"
            stats["docx"] += 1
        else:
            text = _extract_pdf(path)
            kind = "work_order_pdf"
            stats["pdf"] += 1
            if not text.strip():
                stats["pdf_empty_text"] += 1
                text = (
                    f"PDF work order file: {title}. "
                    "Text layer empty or scanned; use RunPulse OCR for full text."
                )
        domains = _guess_domains(path.name, text)
        ministry = "Unknown"
        upper = path.name.upper()
        if "DARPG" in upper:
            ministry = "DARPG / MoPPG related"
        elif "DPIIT" in upper:
            ministry = "DPIIT"
        elif "DOPPW" in upper:
            ministry = "DoPPW"
        _add_doc(
            docs,
            chunks,
            doc_id=doc_id,
            title=title,
            ministry=ministry,
            date=None,
            domains=domains,
            deliverables="",
            full_text=text[:50000],
            value_inr=0.0,
            path=_rel(path),
            kind=kind,
        )
    return stats


@lru_cache(maxsize=1)
def load_corpus() -> tuple[list[CorpusDoc], list[Chunk], dict[str, Any]]:
    docs: list[CorpusDoc] = []
    chunks: list[Chunk] = []
    _load_processed(docs, chunks)
    wo_stats = _load_work_orders(docs, chunks)
    meta = {
        "work_order_extract": wo_stats,
        "processed_dir": str(PROCESSED),
        "work_orders_dir": str(WORK_ORDERS),
    }
    return docs, chunks, meta


def corpus_stats() -> dict[str, Any]:
    docs, chunks, meta = load_corpus()
    by_kind: dict[str, int] = {}
    for d in docs:
        by_kind[d.kind] = by_kind.get(d.kind, 0) + 1
    ministries = sorted({d.ministry for d in docs})
    return {
        "documents": len(docs),
        "chunks": len(chunks),
        "by_kind": by_kind,
        "ministries": ministries[:40],
        "ministry_count": len(ministries),
        **meta,
    }


def search_corpus(
    query: str,
    *,
    limit: int = 8,
    ministry: str | None = None,
    domain: str | None = None,
    kind: str | None = None,
) -> list[dict[str, Any]]:
    q = (query or "").strip()
    if not q:
        return []
    q_tokens = set(_tokens(q)) or {q.lower()}
    _, chunks, _ = load_corpus()
    scored: list[tuple[float, Chunk]] = []
    q_lower = q.lower()

    for ch in chunks:
        if ministry and ministry.lower() not in (ch.ministry or "").lower():
            continue
        if domain and not any(domain.lower() in d.lower() for d in ch.domains):
            continue
        if kind and ch.kind != kind:
            continue
        if not (q_tokens & ch.token_set) and q_lower not in ch.text.lower() and q_lower not in ch.title.lower():
            continue

        score = 0.0
        text_l = ch.text.lower()
        for tok in q_tokens:
            tf = text_l.count(tok)
            if tf:
                score += 1.0 + math.log1p(tf)
            if tok in ch.title.lower():
                score += 1.5
            if any(tok in d.lower() for d in ch.domains):
                score += 0.75
        if q_lower in text_l:
            score += 2.0
        if ch.kind == "processed_json":
            score += 0.35
        elif ch.kind == "work_order_docx":
            score += 0.2
        if "Text layer empty" in ch.text:
            score *= 0.4
        if score > 0:
            scored.append((score, ch))

    scored.sort(key=lambda x: x[0], reverse=True)
    results = []
    for score, ch in scored[: max(1, min(limit, 25))]:
        results.append(
            {
                "chunk_id": ch.chunk_id,
                "doc_id": ch.doc_id,
                "title": ch.title,
                "ministry": ch.ministry,
                "date": ch.date,
                "domains": ch.domains,
                "value_inr": ch.value_inr,
                "text": ch.text,
                "source": ch.source,
                "kind": ch.kind,
                "score": round(score, 4),
            }
        )
    return results


def get_document(doc_id: str) -> dict[str, Any] | None:
    docs, _, _ = load_corpus()
    for d in docs:
        if d.doc_id == doc_id or d.doc_id.replace(" ", "_") == doc_id:
            return {
                "doc_id": d.doc_id,
                "title": d.title,
                "ministry": d.ministry,
                "date": d.date,
                "domains": d.domains,
                "deliverables": d.deliverables,
                "full_text": d.full_text,
                "value_inr": d.value_inr,
                "path": d.path,
                "kind": d.kind,
            }
    return None


def list_documents(limit: int = 300) -> list[dict[str, Any]]:
    docs, _, _ = load_corpus()
    return [
        {
            "doc_id": d.doc_id,
            "title": d.title,
            "ministry": d.ministry,
            "date": d.date,
            "domains": d.domains,
            "value_inr": d.value_inr,
            "path": d.path,
            "kind": d.kind,
        }
        for d in docs[:limit]
    ]


def reload_corpus() -> dict[str, Any]:
    load_corpus.cache_clear()
    return corpus_stats()
